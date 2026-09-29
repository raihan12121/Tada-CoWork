import re
from pathlib import Path
from fastapi import APIRouter, HTTPException, Request
from fastapi.responses import FileResponse, PlainTextResponse
from app.config import settings
from app.sandbox.process_sandbox import sandbox_manager
from app.engine.session_manager import session_manager
from app.core.identity import Principal, require_workspace_access

router = APIRouter(prefix="/artifacts", tags=["artifacts"])

SESSION_ID_PATTERN = re.compile(r"^[a-zA-Z0-9_\-]+$")


async def _assert_artifact_access(request: Request, session_id: str) -> None:
    session = await session_manager.get_session(session_id)
    if not session:
        raise HTTPException(status_code=404, detail="Session not found")
    principal = getattr(request.state, "principal", Principal("anonymous", "default", frozenset({"*"}), frozenset({"local_dev"})))
    try:
        require_workspace_access(principal, session.workspace_id)
    except PermissionError as exc:
        raise HTTPException(status_code=403, detail="Workspace access denied") from exc

@router.get("/download/{session_id}/{filename}")
async def download_artifact(request: Request, session_id: str, filename: str):
    if not SESSION_ID_PATTERN.match(session_id):
        raise HTTPException(status_code=400, detail="Invalid session ID format")
    await _assert_artifact_access(request, session_id)
    sandbox = sandbox_manager.get_existing(session_id)
    safe_name = Path(filename).name
    if ".." in filename or safe_name != filename:
        raise HTTPException(status_code=400, detail="Invalid filename format")
    file_path = (sandbox.artifacts_dir / safe_name) if sandbox else (sandbox_manager.archived_artifacts_dir(session_id) / safe_name)
    
    if not file_path.exists():
        raise HTTPException(status_code=404, detail="Artifact file not found")
        
    return FileResponse(
        path=str(file_path),
        filename=safe_name,
        media_type="application/octet-stream"
    )

@router.get("/preview/{session_id}/{filename}")
async def preview_artifact(request: Request, session_id: str, filename: str):
    if not SESSION_ID_PATTERN.match(session_id):
        raise HTTPException(status_code=400, detail="Invalid session ID format")
    await _assert_artifact_access(request, session_id)
    sandbox = sandbox_manager.get_existing(session_id)
    safe_name = Path(filename).name
    if ".." in filename or safe_name != filename:
        raise HTTPException(status_code=400, detail="Invalid filename format")
    file_path = (sandbox.artifacts_dir / safe_name) if sandbox else (sandbox_manager.archived_artifacts_dir(session_id) / safe_name)
    
    if not file_path.exists():
        raise HTTPException(status_code=404, detail="Artifact file not found")
        
    ext = file_path.suffix.lower()
    try:
        if ext in (".md", ".txt", ".json", ".py", ".js", ".html", ".log"):
            text = file_path.read_text(encoding="utf-8", errors="replace")
            return {"filename": safe_name, "type": ext[1:], "content": text}
        elif ext == ".csv":
            import csv
            with open(file_path, "r", encoding="utf-8", errors="replace") as f:
                reader = csv.reader(f)
                rows = [row for i, row in enumerate(reader) if i < 100]
            text = file_path.read_text(encoding="utf-8", errors="replace")[:10000]
            return {"filename": safe_name, "type": "csv", "rows": rows, "content": text}
        elif ext == ".xlsx":
            import openpyxl
            wb = openpyxl.load_workbook(file_path, data_only=True)
            sheets_data = []
            for sheet_name in wb.sheetnames[:5]:
                ws = wb[sheet_name]
                rows = []
                for row in ws.iter_rows(values_only=True):
                    if any(cell is not None for cell in row):
                        rows.append([str(c) if c is not None else "" for c in row[:25]])
                    if len(rows) >= 100:
                        break
                sheets_data.append({"name": sheet_name, "rows": rows})
            return {
                "filename": safe_name,
                "type": "xlsx",
                "sheets": sheets_data,
                "content": f"Spreadsheet with {len(wb.sheetnames)} sheet(s)."
            }
        elif ext == ".docx":
            import docx
            doc = docx.Document(file_path)
            sections = []
            for p in doc.paragraphs:
                text = p.text.strip()
                if text:
                    style = p.style.name if p.style else ""
                    sections.append({"text": text, "is_heading": "heading" in style.lower(), "style": style})
            tables_data = []
            for t in doc.tables[:5]:
                t_rows = []
                for row in t.rows[:25]:
                    t_rows.append([c.text.strip() for c in row.cells])
                tables_data.append(t_rows)
            return {
                "filename": safe_name,
                "type": "docx",
                "sections": sections[:60],
                "tables": tables_data,
                "content": "\n\n".join(s["text"] for s in sections[:40])
            }
        elif ext == ".pptx":
            from pptx import Presentation
            prs = Presentation(file_path)
            slides_data = []
            for i, slide in enumerate(prs.slides, 1):
                if i > 15:
                    break
                title = slide.shapes.title.text if slide.shapes.title else f"Slide {i}"
                bullets = []
                for shape in slide.shapes:
                    if shape.has_text_frame and shape != slide.shapes.title:
                        for p in shape.text_frame.paragraphs:
                            if p.text.strip():
                                bullets.append(p.text.strip())
                slides_data.append({"index": i, "title": title, "content": bullets})
            return {
                "filename": safe_name,
                "type": "pptx",
                "slides": slides_data,
                "content": "\n".join(f"Slide {s['index']}: {s['title']}\n" + "\n".join(f" - {b}" for b in s['content']) for s in slides_data)
            }
        elif ext in (".png", ".jpg", ".jpeg", ".webp", ".svg", ".gif"):
            return {
                "filename": safe_name,
                "type": "image",
                "download_url": f"/v1/artifacts/download/{session_id}/{safe_name}",
                "content": f"[Image deliverable: {safe_name}]"
            }
        elif ext == ".pdf":
            return {
                "filename": safe_name,
                "type": "pdf",
                "file_size_bytes": file_path.stat().st_size,
                "download_url": f"/v1/artifacts/download/{session_id}/{safe_name}",
                "content": f"[PDF Document ({file_path.stat().st_size / 1024:.1f} KB)]"
            }
        else:
            return {
                "filename": safe_name,
                "type": ext[1:],
                "content": f"[Binary deliverable: {ext.upper()} document ({file_path.stat().st_size} bytes). Click Download to open.]"
            }
    except Exception as exc:
        return {
            "filename": safe_name,
            "type": ext[1:],
            "content": f"Preview parsing failed: {exc}. Please click Download to open the deliverable."
        }
